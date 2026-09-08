from __future__ import annotations

import csv
import io
from typing import Iterable


FIELDS = ("id", "product", "price", "unit", "currency", "source", "source_date", "collected_at", "variation", "variation_pct", "is_unchanged", "notes")


def csv_bytes(rows: Iterable[dict]) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=FIELDS, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


def xlsx_bytes(rows: list[dict]) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("Installez openpyxl pour activer l'export Excel (.xlsx).") from exc
    output = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Relevés"
    headers = ["ID", "Produit", "Prix", "Unité", "Devise", "Source", "Date source", "Date collecte", "Variation", "Variation %", "Inchangée", "Remarque"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B2E4F")
    for row in rows:
        sheet.append([row.get(field, "") for field in FIELDS])
    sheet.freeze_panes = "A2"
    for column, width in {"A": 8, "B": 16, "C": 14, "D": 16, "F": 34, "G": 14, "H": 24, "L": 46}.items():
        sheet.column_dimensions[column].width = width
    workbook.save(output)
    return output.getvalue()

